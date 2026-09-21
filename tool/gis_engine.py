import os
import sys
import zipfile
import time
from lxml import etree
import shapely.geometry as sg
from shapely.ops import unary_union
from shapely.prepared import prep
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

def parse_coordinates(coord_text):
    """Parse KML coordinate text into list of (lon, lat) tuples."""
    if not coord_text:
        return []
    coords = []
    tokens = coord_text.strip().split()
    for token in tokens:
        parts = token.split(',')
        if len(parts) >= 2:
            try:
                coords.append((float(parts[0]), float(parts[1])))
            except ValueError:
                continue
    return coords

def dd_to_dms(deg, is_lat=True):
    """Convert decimal degrees to Degrees Minutes Seconds (DMS) string."""
    direction = 'N' if is_lat else 'E'
    if deg < 0:
        direction = 'S' if is_lat else 'W'
        deg = abs(deg)
    d = int(deg)
    m = int((deg - d) * 60)
    s = (deg - d - m / 60) * 3600
    return f"{d}°{m:02d}'{s:05.2f}\"{direction}"

def format_coord_dd_dms(lon, lat):
    """Format coordinate as combined DD and DMS string."""
    lat_dms = dd_to_dms(lat, is_lat=True)
    lon_dms = dd_to_dms(lon, is_lat=False)
    return f"{lat:.6f}, {lon:.6f} ({lat_dms}, {lon_dms})"

def extract_geometries_from_elem(elem):
    """Extract Shapely geometry from an XML element (Placemark)."""
    geoms = []
    
    # 1. Point
    points = elem.xpath('.//*[local-name()="Point"]/*[local-name()="coordinates"]/text()')
    for pt_text in points:
        coords = parse_coordinates(pt_text)
        if coords:
            geoms.append(sg.Point(coords[0]))
            
    # 2. LineString
    linestrings = elem.xpath('.//*[local-name()="LineString"]/*[local-name()="coordinates"]/text()')
    for ls_text in linestrings:
        coords = parse_coordinates(ls_text)
        if len(coords) >= 2:
            geoms.append(sg.LineString(coords))
            
    # 3. LinearRing
    linear_rings = elem.xpath('.//*[local-name()="LinearRing"]/*[local-name()="coordinates"]/text()')
    for lr_text in linear_rings:
        coords = parse_coordinates(lr_text)
        if len(coords) >= 3:
            geoms.append(sg.LinearRing(coords))

    # 4. Polygon
    for poly in elem.xpath('.//*[local-name()="Polygon"]'):
        outer = poly.xpath('.//*[local-name()="outerBoundaryIs"]//*[local-name()="coordinates"]/text()')
        if outer:
            ext_coords = parse_coordinates(outer[0])
            if len(ext_coords) >= 3:
                inners = []
                for inner in poly.xpath('.//*[local-name()="innerBoundaryIs"]//*[local-name()="coordinates"]/text()'):
                    in_coords = parse_coordinates(inner)
                    if len(in_coords) >= 3:
                        inners.append(in_coords)
                try:
                    poly_geom = sg.Polygon(ext_coords, inners)
                    if not poly_geom.is_valid:
                        poly_geom = poly_geom.buffer(0)
                    geoms.append(poly_geom)
                except Exception:
                    pass
    
    if not geoms:
        return None
    elif len(geoms) == 1:
        return geoms[0]
    else:
        return sg.GeometryCollection(geoms)

def read_kmz_or_kml_features(file_path, progress_callback=None):
    """
    Read all Placemarks from a KMZ or KML file.
    Returns a list of dicts with: name, description, geom_type, geometry, attributes.
    """
    features = []
    is_kmz = file_path.lower().endswith('.kmz')
    
    if is_kmz:
        with zipfile.ZipFile(file_path, 'r') as z:
            kml_names = [n for n in z.namelist() if n.lower().endswith('.kml')]
            if not kml_names:
                raise ValueError("Không tìm thấy file .kml bên trong file .kmz")
            
            for kml_name in kml_names:
                kml_bytes = z.read(kml_name)
                root = etree.fromstring(kml_bytes)
                placemarks = root.xpath('//*[local-name()="Placemark"]')
                total = len(placemarks)
                for idx, pm in enumerate(placemarks):
                    name_el = pm.xpath('./*[local-name()="name"]/text()')
                    desc_el = pm.xpath('./*[local-name()="description"]/text()')
                    name = name_el[0].strip() if name_el else f"Đối tượng {idx+1}"
                    desc = desc_el[0].strip() if desc_el else ""
                    
                    attrs = {}
                    for d in pm.xpath('./*[local-name()="ExtendedData"]/*[local-name()="Data"]'):
                        dname = d.get('name')
                        dval = d.xpath('./*[local-name()="value"]/text()')
                        attrs[dname] = dval[0] if dval else ''
                    for sd in pm.xpath('./*[local-name()="ExtendedData"]/*[local-name()="SchemaData"]/*[local-name()="SimpleData"]'):
                        sdname = sd.get('name')
                        attrs[sdname] = sd.text if sd.text else ''
                        
                    geom = extract_geometries_from_elem(pm)
                    if geom and not geom.is_empty:
                        features.append({
                            'index': idx + 1,
                            'name': name,
                            'description': desc,
                            'geom_type': geom.geom_type,
                            'geometry': geom,
                            'attributes': attrs
                        })
                    if progress_callback and total > 0 and idx % 50 == 0:
                        progress_callback(f"Đọc file nghiên cứu: {idx+1}/{total} đối tượng", int((idx+1)/total * 30))
    else:
        context = etree.iterparse(file_path, events=('end',), tag='{http://www.opengis.net/kml/2.2}Placemark')
        idx = 0
        for event, pm in context:
            idx += 1
            name_el = pm.xpath('./*[local-name()="name"]/text()')
            desc_el = pm.xpath('./*[local-name()="description"]/text()')
            name = name_el[0].strip() if name_el else f"Đối tượng {idx}"
            desc = desc_el[0].strip() if desc_el else ""
            
            attrs = {}
            for d in pm.xpath('./*[local-name()="ExtendedData"]/*[local-name()="Data"]'):
                dname = d.get('name')
                dval = d.xpath('./*[local-name()="value"]/text()')
                attrs[dname] = dval[0] if dval else ''
                
            geom = extract_geometries_from_elem(pm)
            if geom and not geom.is_empty:
                features.append({
                    'index': idx,
                    'name': name,
                    'description': desc,
                    'geom_type': geom.geom_type,
                    'geometry': geom,
                    'attributes': attrs
                })
            pm.clear()
            while pm.getprevious() is not None:
                del pm.getparent()[0]
                
    return features

def analyze_intersection(study_file, boundary_file, progress_callback=None):
    """
    Main function to analyze overlap between study area and administrative boundaries.
    Calculates center coordinates (Centroid, Bounding Box Center, Representative Point).
    """
    if progress_callback:
        progress_callback("Đang đọc file ranh nghiên cứu...", 5)
        
    study_features = read_kmz_or_kml_features(study_file, progress_callback)
    if not study_features:
        raise ValueError("Không tìm thấy dữ liệu hình học hợp lệ trong file ranh nghiên cứu.")
        
    all_study_geoms = [f['geometry'] for f in study_features]
    study_union = unary_union(all_study_geoms)
    if not study_union.is_valid:
        study_union = study_union.buffer(0)
    minx, miny, maxx, maxy = study_union.bounds
    study_box = sg.box(minx, miny, maxx, maxy)
    
    # 1. Calculate Center Coordinates
    centroid = study_union.centroid
    centroid_lon = centroid.x
    centroid_lat = centroid.y
    
    bbox_center_lon = (minx + maxx) / 2.0
    bbox_center_lat = (miny + maxy) / 2.0
    
    rep_point = study_union.representative_point()
    rep_lon = rep_point.x
    rep_lat = rep_point.y
    
    gmaps_url = f"https://www.google.com/maps?q={centroid_lat:.6f},{centroid_lon:.6f}"
    
    prep_study_union = prep(study_union)
    
    if progress_callback:
        progress_callback(f"Đã nạp {len(study_features)} đối tượng dự án. Đang quét ranh xã phường...", 35)
        
    matched_communes = []
    is_boundary_kmz = boundary_file.lower().endswith('.kmz')
    
    def process_commune_placemark(elem, idx):
        name_el = elem.xpath('./*[local-name()="name"]/text()')
        desc_el = elem.xpath('./*[local-name()="description"]/text()')
        name = name_el[0].strip() if name_el else f"Xã/Phường {idx}"
        desc = desc_el[0].strip() if desc_el else ""
        
        attrs = {}
        for d in elem.xpath('./*[local-name()="ExtendedData"]/*[local-name()="Data"]'):
            dname = d.get('name')
            dval = d.xpath('./*[local-name()="value"]/text()')
            attrs[dname] = dval[0] if dval else ''
        for sd in elem.xpath('./*[local-name()="ExtendedData"]/*[local-name()="SchemaData"]/*[local-name()="SimpleData"]'):
            sdname = sd.get('name')
            attrs[sdname] = sd.text if sd.text else ''
            
        geom = extract_geometries_from_elem(elem)
        if geom and not geom.is_empty:
            # 1. Fast bounding-box check
            if study_box.intersects(geom.envelope):
                # 2. Prepared geometry intersection check
                if prep_study_union.intersects(geom):
                    prep_geom = prep(geom)
                    contained_study_features = []
                    for sf in study_features:
                        if prep_geom.intersects(sf['geometry']):
                            contained_study_features.append(sf['name'])
                    
                    commune_centroid = geom.centroid
                    matched_communes.append({
                        'ma_xa': attrs.get('ma_xa', ''),
                        'ten_xa': attrs.get('ten_xa', name),
                        'loai': attrs.get('loai', ''),
                        'cap': attrs.get('cap', ''),
                        'stt': attrs.get('stt', ''),
                        'ma_tinh': attrs.get('ma_tinh', ''),
                        'ten_tinh': attrs.get('ten_tinh', ''),
                        'dtich_km2': attrs.get('dtich_km2', ''),
                        'dan_so': attrs.get('dan_so', ''),
                        'matdo_km2': attrs.get('matdo_km2', ''),
                        'sap_nhap': attrs.get('sap_nhap', ''),
                        'tru_so': attrs.get('tru_so', ''),
                        'so_doi_tuong_giao': len(contained_study_features),
                        'cac_doi_tuong_giao': ", ".join(contained_study_features[:30]) + (f" (+{len(contained_study_features)-30} khác)" if len(contained_study_features) > 30 else ""),
                        'toa_do_tam_xa': f"{commune_centroid.x:.6f}, {commune_centroid.y:.6f}",
                        'geometry': geom
                    })
    
    if is_boundary_kmz:
        with zipfile.ZipFile(boundary_file, 'r') as z:
            kml_names = [n for n in z.namelist() if n.lower().endswith('.kml')]
            for kml_name in kml_names:
                kml_bytes = z.read(kml_name)
                root = etree.fromstring(kml_bytes)
                placemarks = root.xpath('//*[local-name()="Placemark"]')
                total = len(placemarks)
                for idx, pm in enumerate(placemarks):
                    process_commune_placemark(pm, idx + 1)
                    if progress_callback and idx % 100 == 0:
                        progress_callback(f"Đang duyệt ranh xã: {idx+1}/{total}", 35 + int((idx+1)/total * 50))
    else:
        context = etree.iterparse(boundary_file, events=('end',), tag='{http://www.opengis.net/kml/2.2}Placemark')
        idx = 0
        for event, elem in context:
            idx += 1
            process_commune_placemark(elem, idx)
            if progress_callback and idx % 200 == 0:
                progress_callback(f"Đang quét ranh xã: đã duyệt {idx} xã/phường...", min(85, 35 + int(idx / 3500 * 50)))
            elem.clear()
            while elem.getprevious() is not None:
                del elem.getparent()[0]

    if progress_callback:
        progress_callback("Phân tích chi tiết từng đối tượng dự án...", 90)

    # Determine which commune contains the study area centroid
    center_commune_name = "Nằm trên ranh giới các xã"
    center_point = sg.Point(centroid_lon, centroid_lat)
    for c in matched_communes:
        if c['geometry'].contains(center_point) or c['geometry'].intersects(center_point.buffer(0.0001)):
            center_commune_name = f"{c['ten_xa']} ({c['ten_tinh']})"
            break

    # Breakdown per study feature using prepared commune geometries
    prepared_communes = [(c, prep(c['geometry'])) for c in matched_communes]
    feature_breakdown = []
    
    for sf in study_features:
        matched_commune_names = []
        matched_province_names = []
        for c, p_c in prepared_communes:
            if p_c.intersects(sf['geometry']):
                matched_commune_names.append(c['ten_xa'])
                if c['ten_tinh'] and c['ten_tinh'] not in matched_province_names:
                    matched_province_names.append(c['ten_tinh'])
                    
        f_centroid = sf['geometry'].centroid
        feature_breakdown.append({
            'STT': sf['index'],
            'Tên đối tượng': sf['name'],
            'Loại hình học': sf['geom_type'],
            'Xã / Phường': ", ".join(matched_commune_names) if matched_commune_names else "Ngoài ranh khớp",
            'Tỉnh / Thành': ", ".join(matched_province_names) if matched_province_names else "",
            'Tọa độ tâm (Kinh độ, Vĩ độ)': f"{f_centroid.x:.6f}, {f_centroid.y:.6f}",
            'Tọa độ tâm (Độ Phút Giây)': f"{dd_to_dms(f_centroid.y, True)}, {dd_to_dms(f_centroid.x, False)}",
            'Mô tả': sf['description']
        })

    if progress_callback:
        progress_callback(f"Hoàn thành phân tích! Khớp {len(matched_communes)} xã/phường.", 100)

    center_info = {
        'centroid_lon': centroid_lon,
        'centroid_lat': centroid_lat,
        'centroid_dd_str': f"{centroid_lat:.6f}, {centroid_lon:.6f}",
        'centroid_dms_str': f"{dd_to_dms(centroid_lat, True)}, {dd_to_dms(centroid_lon, False)}",
        'centroid_full_str': format_coord_dd_dms(centroid_lon, centroid_lat),
        'bbox_center_lon': bbox_center_lon,
        'bbox_center_lat': bbox_center_lat,
        'bbox_center_str': format_coord_dd_dms(bbox_center_lon, bbox_center_lat),
        'rep_point_lon': rep_lon,
        'rep_point_lat': rep_lat,
        'rep_point_str': format_coord_dd_dms(rep_lon, rep_lat),
        'center_commune': center_commune_name,
        'google_maps_url': gmaps_url
    }

    return {
        'study_features': study_features,
        'study_union': study_union,
        'bounds': (minx, miny, maxx, maxy),
        'center_info': center_info,
        'matched_communes': matched_communes,
        'feature_breakdown': feature_breakdown
    }

def export_results_to_excel(results, output_path, study_file_name="", boundary_file_name=""):
    """
    Exports the analysis results to a beautifully formatted Excel file (.xlsx)
    including comprehensive center coordinates and Google Maps link.
    """
    communes_data = []
    for i, c in enumerate(results['matched_communes'], start=1):
        communes_data.append({
            'STT': i,
            'Mã Xã': c.get('ma_xa', ''),
            'Tên Xã / Phường': c.get('ten_xa', ''),
            'Loại': c.get('loai', ''),
            'Cấp': c.get('cap', ''),
            'Mã Tỉnh': c.get('ma_tinh', ''),
            'Tên Tỉnh / Thành': c.get('ten_tinh', ''),
            'Tọa độ tâm xã (Kinh độ, Vĩ độ)': c.get('toa_do_tam_xa', ''),
            'Diện tích (km²)': float(c['dtich_km2']) if c.get('dtich_km2') and c['dtich_km2'].replace('.', '', 1).isdigit() else c.get('dtich_km2', ''),
            'Dân số (người)': int(float(c['dan_so'])) if c.get('dan_so') and c['dan_so'].replace('.', '', 1).isdigit() else c.get('dan_so', ''),
            'Mật độ (người/km²)': float(c['matdo_km2']) if c.get('matdo_km2') and c['matdo_km2'].replace('.', '', 1).isdigit() else c.get('matdo_km2', ''),
            'Sáp nhập': c.get('sap_nhap', ''),
            'Trụ sở': c.get('tru_so', ''),
            'Số đối tượng giao cắt': c.get('so_doi_tuong_giao', 0),
            'Danh sách đối tượng giao cắt': c.get('cac_doi_tuong_giao', '')
        })
    df_communes = pd.DataFrame(communes_data)
    df_features = pd.DataFrame(results['feature_breakdown'])

    minx, miny, maxx, maxy = results['bounds']
    center = results.get('center_info', {})
    
    metadata = [
        {'Chỉ tiêu': 'Tên file ranh nghiên cứu / dự án', 'Giá trị': study_file_name or 'N/A'},
        {'Chỉ tiêu': 'Tên file ranh hành chính', 'Giá trị': boundary_file_name or 'N/A'},
        {'Chỉ tiêu': 'Thời gian xuất báo cáo', 'Giá trị': time.strftime("%Y-%m-%d %H:%M:%S")},
        {'Chỉ tiêu': 'Tổng số đối tượng trong file nghiên cứu', 'Giá trị': len(results['study_features'])},
        {'Chỉ tiêu': 'Tổng số Xã / Phường giao cắt', 'Giá trị': len(results['matched_communes'])},
        {'Chỉ tiêu': '----------------------------------------', 'Giá trị': '----------------------------------------'},
        {'Chỉ tiêu': '📍 TỌA ĐỘ TÂM RANH NGHIÊN CỨU (CENTROID - DD)', 'Giá trị': f"Vĩ độ: {center.get('centroid_lat', 0):.6f}, Kinh độ: {center.get('centroid_lon', 0):.6f}"},
        {'Chỉ tiêu': '📍 TỌA ĐỘ TÂM RANH NGHIÊN CỨU (DMS)', 'Giá trị': center.get('centroid_dms_str', '')},
        {'Chỉ tiêu': '🏛️ Xã / Phường chứa tọa độ tâm dự án', 'Giá trị': center.get('center_commune', '')},
        {'Chỉ tiêu': '📍 Tọa độ tâm hình hộp bao (Bounding Box Center)', 'Giá trị': center.get('bbox_center_str', '')},
        {'Chỉ tiêu': '📍 Tọa độ điểm đại diện bên trong ranh (Representative Point)', 'Giá trị': center.get('rep_point_str', '')},
        {'Chỉ tiêu': '🌐 Link xem vị trí tâm trên Google Maps', 'Giá trị': center.get('google_maps_url', '')},
        {'Chỉ tiêu': '----------------------------------------', 'Giá trị': '----------------------------------------'},
        {'Chỉ tiêu': 'Phạm vi kinh độ (Min Lon - Max Lon)', 'Giá trị': f"{minx:.6f} - {maxx:.6f}"},
        {'Chỉ tiêu': 'Phạm vi vĩ độ (Min Lat - Max Lat)', 'Giá trị': f"{miny:.6f} - {maxy:.6f}"},
    ]
    df_meta = pd.DataFrame(metadata)

    # Write sheets with permission fallback if file is open in Excel
    actual_output_path = output_path
    try:
        with pd.ExcelWriter(actual_output_path, engine='openpyxl') as writer:
            df_communes.to_excel(writer, sheet_name='Xã Phường Giao Cắt', index=False)
            df_features.to_excel(writer, sheet_name='Chi Tiết Đối Tượng', index=False)
            df_meta.to_excel(writer, sheet_name='Thông Tin Tổng Quan', index=False)
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        actual_output_path = f"{base}_{time.strftime('%Y%m%d_%H%M%S')}{ext}"
        with pd.ExcelWriter(actual_output_path, engine='openpyxl') as writer:
            df_communes.to_excel(writer, sheet_name='Xã Phường Giao Cắt', index=False)
            df_features.to_excel(writer, sheet_name='Chi Tiết Đối Tượng', index=False)
            df_meta.to_excel(writer, sheet_name='Thông Tin Tổng Quan', index=False)

    # Styling Excel workbook
    wb = openpyxl.load_workbook(actual_output_path)
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
    highlight_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    highlight_font = Font(name="Segoe UI", size=10, bold=True, color="92400E")
    regular_font = Font(name="Segoe UI", size=10)
    link_font = Font(name="Segoe UI", size=10, color="0000FF", underline="single")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        
        # Header formatting
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data rows formatting
        for row_idx in range(2, ws.max_row + 1):
            is_even = (row_idx % 2 == 0)
            ws.row_dimensions[row_idx].height = 22
            col1_val = str(ws.cell(row=row_idx, column=1).value or '')
            is_center_row = 'TỌA ĐỘ TÂM' in col1_val or 'Google Maps' in col1_val or 'Xã / Phường chứa' in col1_val

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                
                if is_center_row:
                    cell.fill = highlight_fill
                    cell.font = highlight_font
                elif is_even:
                    cell.fill = zebra_fill
                    cell.font = regular_font
                else:
                    cell.font = regular_font

                # Hyperlink if it is a URL
                val_str = str(cell.value or '')
                if val_str.startswith('http://') or val_str.startswith('https://'):
                    cell.hyperlink = val_str
                    cell.font = link_font

                # Alignment
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    
        # Adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                lines = val.split('\n')
                for l in lines:
                    if len(l) > max_len:
                        max_len = len(l)
            adjusted_width = min(max(max_len + 4, 14), 70)
            ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(actual_output_path)
    return actual_output_path
